"""Diccionario de nombres alternativos de municipios (código INE → variantes) a partir de las fuentes que traen código y nombre:
Hacienda (deuda viva), Seguridad Social (MUNCNAE) y SEPE (ficheros con código). Lo usa geo.py para casar fuentes que solo traen nombre.
→ cache/nombres_alias.json"""
import glob, json, os, sys
import openpyxl, xlrd
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from geo import norm
HERE = os.path.dirname(os.path.abspath(__file__)); A = {}
def add(code, name):
    if code and name and str(name).strip(): A.setdefault(code, set()).add(norm(name))
for f in glob.glob(f"{HERE}/cache/externos/hacienda_deuda/*.xls*"):
    if f.endswith(".xls"):
        sh = xlrd.open_workbook(f).sheet_by_index(0); rows = [sh.row_values(r) for r in range(sh.nrows)]
    else: rows = [list(r) for r in openpyxl.load_workbook(f, read_only=True).worksheets[0].iter_rows(values_only=True)]
    for r in rows:
        try: add(f"{int(float(r[3])):02d}{int(float(r[5])):03d}", r[6])
        except (TypeError, ValueError, IndexError): pass
for p in glob.glob(f"{HERE}/cache/externos/sepe_parsed6/*.json"):
    for c, (pv, nm) in json.load(open(p)).get("names", {}).items(): A.setdefault(c, set()).add(nm)
for f in [f"{HERE}/cache/externos/ss_muncnae/MUNCNAE1225.xlsx", f"{HERE}/cache/externos/ss_muncnae_hist/MUNCNAE1211.xlsx"]:
    seen = set()
    for i, r in enumerate(openpyxl.load_workbook(f, read_only=True).worksheets[0].iter_rows(values_only=True)):
        if i == 0 or r[3] in seen: continue
        seen.add(r[3])
        try: add(f"{int(r[3]):05d}", r[4])
        except (TypeError, ValueError): pass
json.dump({c: sorted(v) for c, v in A.items()}, open(f"{HERE}/cache/nombres_alias.json", "w"), ensure_ascii=False)
print(len(A), A.get("07040"), A.get("12135"), A.get("07054"))
