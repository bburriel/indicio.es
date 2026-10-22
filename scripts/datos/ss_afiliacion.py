"""Afiliación a la Seguridad Social por municipio (último día de septiembre): afiliados por habitante, peso de cada régimen y cambio.
Excel mensual «Datos Municipios» de la Seguridad Social (https://www.seg-social.es/descarga/es/MuniMMAAAA → redirige al .xlsx).
Ojo: la afiliación se asigna al municipio de la cuenta de cotización (donde trabaja o tiene la sede la empresa), no al de residencia,
por eso mide empleo situado en el municipio y puede superar los 100 afiliados por cada 100 habitantes.
Uso: python3 ss_afiliacion.py   (descarga a cache/ss/ y escribe cache/out/afiliacion.json)"""
import json, os, re, subprocess, warnings
import openpyxl
warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__)); os.makedirs(f"{HERE}/cache/ss", exist_ok=True); os.makedirs(f"{HERE}/cache/out", exist_ok=True)
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"
def bajar(mes):
    f = f"{HERE}/cache/ss/Muni{mes}.xlsx"
    if not os.path.exists(f) or os.path.getsize(f) < 1000:
        subprocess.run(["curl", "-sS", "-m", "180", "-A", UA, "-L", "-o", f, f"https://www.seg-social.es/descarga/es/Muni{mes}"], check=True)
    return f
def num(v):
    """'4.010' → 4010 · '>=768' → 768 · ' <5' → 2 (valor oculto por secreto estadístico) · '' → None"""
    if v is None: return None
    s = str(v).strip().replace(">=", "").replace(".", "")
    if s.startswith("<"): return 2.0
    try: return float(s)
    except ValueError: return None
def leer(mes):
    ws = openpyxl.load_workbook(bajar(mes), data_only=True).worksheets[0]
    cab = [str(c).strip() if c else "" for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    out = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        m = re.match(r"(\d{5}) ", str(row[1] or ""))
        if m: out[m.group(1)] = {cab[i]: num(row[i]) for i in range(3, len(cab)) if cab[i]}
    return out
a26, a21, a19 = leer("092026"), leer("092021"), leer("092019")
ix = json.load(open(f"{HERE}/../../data/index.json")); POP = {r[0]: r[ix["fields"].index("pop")] for r in ix["rows"]}
def pc(a, b): return round(a / b * 100, 1) if a is not None and b else None
datos = {}
for code, d in a26.items():
    pop, t = POP.get(code), d.get("TOTAL")
    if not pop or pop < 1000 or not t or t < 20: continue
    o = {"afilHab": pc(t, pop), "afilAutonomos": pc(d.get("R. E. T. Autónomos"), t), "afilAgrario": pc(d.get("R. G.- S.E.Agrario"), t), "afilHogar": pc(d.get("R. G.- S.E.Hogar"), t)}
    for clave, base in (("afilVar", a21), ("afilVar19", a19)):
        b = (base.get(code) or {}).get("TOTAL")
        if b and b >= 100: o[clave] = round((t / b - 1) * 100, 1)
    datos[code] = {k: v for k, v in o.items() if v is not None}
SRC = "Seguridad Social, afiliados por municipio, septiembre de 2026"
INDICADORES = [
 ("afilHab", "Afiliados a la Seguridad Social por cada 100 habitantes", "Trabajo", "pct", SRC + " (empleo situado en el municipio, no residentes: puede pasar de 100)"),
 ("afilAutonomos", "Afiliados que son autónomos", "Trabajo", "pct", SRC + " (% de los afiliados)"),
 ("afilAgrario", "Afiliados del régimen agrario", "Trabajo", "pct", SRC + " (% de los afiliados; sistema especial agrario)"),
 ("afilHogar", "Afiliadas al servicio doméstico", "Trabajo", "pct", SRC + " (% de los afiliados; sistema especial de empleados del hogar)"),
 ("afilVar", "Cambio del número de afiliados desde 2021", "Trabajo", "pct_chg", "Seguridad Social, afiliados por municipio, septiembre de 2021 a septiembre de 2026"),
 ("afilVar19", "Cambio del número de afiliados desde 2019", "Trabajo", "pct_chg", "Seguridad Social, afiliados por municipio, septiembre de 2019 a septiembre de 2026"),
]
json.dump({"indicadores": INDICADORES, "datos": datos}, open(f"{HERE}/cache/out/afiliacion.json", "w"), ensure_ascii=False)
print(len(datos), "municipios")
