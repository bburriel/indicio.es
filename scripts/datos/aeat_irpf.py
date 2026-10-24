"""IRPF por municipios (Agencia Tributaria, «Estadística de los declarantes del IRPF por municipios»).
Lee los Excel por comunidad autónoma que publica la AEAT (territorio fiscal común: sin Navarra ni País Vasco)
y calcula, por municipio de más de 1.000 habitantes con dato no sujeto a secreto estadístico, la renta
bruta y disponible media por declarante y la composición de la renta bruta.
Uso: python3 aeat_irpf.py  → cache/out/irpf.json"""
import json, glob, os, re, warnings
import openpyxl
warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__)); os.makedirs(f"{HERE}/cache/out", exist_ok=True)
AÑO = 2024
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None
def hoja(ws):
    hdr = next((r for r in range(1, 80) if ws.cell(r, 8).value and "digo" in str(ws.cell(r, 8).value)), None)
    cols = {c: re.sub(r"\s+", " ", str(ws.cell(hdr, c).value)).strip() for c in range(9, ws.max_column + 1) if ws.cell(hdr, c).value}
    out = {}
    for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
        code = row[7] if len(row) > 7 else None
        if code and re.fullmatch(r"\d{5}", str(code).strip()):
            out[str(code).strip()] = {cols[c]: num(row[c - 1]) for c in cols if c - 1 < len(row)}
    return out
res = {}
for f in sorted(glob.glob(f"{HERE}/cache/aeat/*.xlsx")):
    wb = openpyxl.load_workbook(f, data_only=True)
    n, imp = hoja(wb["1.1"]), hoja(wb["1.4"])
    for code, d in n.items():
        res.setdefault(code, {})["n"] = d.get("Número total declaraciones")
    for code, d in imp.items():
        res.setdefault(code, {}).update({"bruta": d.get("RENTA BRUTA"), "disp": d.get("RENTA DISPONIBLE"), "trabajo": d.get("RENTAS DEL TRABAJO"),
            "mobil": d.get("RENTAS DEL CAPITAL MOBILIARIO"), "inmob": d.get("RENTAS DE BIENES INMUEBLES NO AFECTOS A ACTIV. ECONÓMICAS"),
            "activ": d.get("RENTAS DE ACTIVIDADES ECONÓMICAS"), "gan": d.get("GANANCIAS PATRIMONIALES NETAS"), "exenta": d.get("RENTAS EXENTAS")})
    print(os.path.basename(f), len(n), "municipios")
json.dump(res, open(f"{HERE}/cache/out/irpf_raw.json", "w"))
print(len(res), "municipios en total")

# ---- indicadores por municipio (importes en euros; se descartan municipios con menos de 100 declaraciones) ----
ix = json.load(open(f"{HERE}/../../data/index.json")); POP = {r[0]: r[ix["fields"].index("pop")] for r in ix["rows"]}
def pc(a, b): return round(a / b * 100, 1) if a is not None and b else None
datos = {}
for code, d in res.items():
    n, b, dp = d.get("n"), d.get("bruta"), d.get("disp")
    if code not in POP or not n or n < 100 or not b: continue
    o = {"irpfBruta": round(b / n), "irpfDecl": pc(n, POP[code])}
    if dp: o["irpfDisp"] = round(dp / n); o["irpfCarga"] = pc(b - dp, b)
    o["irpfTrabajo"] = pc(d.get("trabajo"), b); o["irpfActiv"] = pc(d.get("activ"), b); o["irpfExenta"] = pc(d.get("exenta"), b)
    cap = [d.get(k) for k in ("mobil", "inmob", "gan")]
    if all(c is not None for c in cap): o["irpfCapital"] = pc(sum(cap), b)
    datos[code] = {k: v for k, v in o.items() if v is not None}
SRC = f"Agencia Tributaria, Estadística de los declarantes del IRPF por municipios, ejercicio {AÑO}"
INDICADORES = [
 ("irpfBruta", "Renta bruta media por declarante (IRPF)", "Riqueza", "eur", SRC + " (territorio fiscal común)"),
 ("irpfDisp", "Renta disponible media por declarante (IRPF)", "Riqueza", "eur", SRC + " (tras impuestos y cotizaciones)"),
 ("irpfCarga", "Impuestos y cotizaciones sobre la renta bruta", "Riqueza", "pct", SRC + " (cuota líquida y cotizaciones, % de la renta bruta)"),
 ("irpfTrabajo", "Renta que viene del trabajo y las pensiones", "Riqueza", "pct", SRC + " (% de la renta bruta)"),
 ("irpfActiv", "Renta que viene de actividades económicas", "Riqueza", "pct", SRC + " (% de la renta bruta; autónomos y empresarios)"),
 ("irpfCapital", "Renta del capital y el patrimonio", "Riqueza", "pct", SRC + " (% de la renta bruta; intereses, alquileres y ganancias)"),
 ("irpfExenta", "Renta exenta de impuestos", "Riqueza", "pct", SRC + " (% de la renta bruta)"),
 ("irpfDecl", "Declaraciones de IRPF por cada 100 habitantes", "Riqueza", "pct", SRC),
]
if not datos:   # sin los Excel de la AEAT no se pisa el resultado anterior (construir.py borraría los indicadores)
    raise SystemExit("No hay datos de IRPF: faltan los Excel en cache/aeat/. No se toca cache/out/irpf.json.")
json.dump({"indicadores": INDICADORES, "datos": datos}, open(f"{HERE}/cache/out/irpf.json", "w"), ensure_ascii=False)
print(len(datos), "municipios con indicadores de IRPF")
